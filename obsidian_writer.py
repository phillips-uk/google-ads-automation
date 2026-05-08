"""Helpers for writing Obsidian-compatible markdown reports and branded xlsx files."""

import os
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Brand colours ───────────────────────────────────────────────────────────
AMBER      = "985830"
WHITE      = "FFFFFF"
ROW_ALT    = "FDF6EE"
DARK_TEXT  = "1A1A1A"
BORDER_CLR = "E8D8C4"

_HEADER_FILL    = PatternFill("solid", fgColor=AMBER)
_BODY_FONT      = Font(name="Arial", size=10, color=DARK_TEXT)
_HEADER_FONT    = Font(bold=True, color=WHITE, name="Arial", size=10)
_TITLE_FONT     = Font(name="Arial", size=11, color=DARK_TEXT, bold=True)
_BOLD_DARK_FONT = Font(bold=True, color=DARK_TEXT, name="Arial", size=10)
_ALT_FILL       = PatternFill("solid", fgColor=ROW_ALT)
_WHITE_FILL     = PatternFill("solid", fgColor=WHITE)
_NO_BORDER      = Border()
_BLACK_MEDIUM   = Side(style="medium", color="000000")
_CF_GREEN       = Font(name="Arial", size=10, color="217346", bold=True)  # conditional: good result
_CF_RED         = Font(name="Arial", size=10, color="C00000", bold=True)  # conditional: bad result
_CENTER         = Alignment(horizontal="center", vertical="center")
_LEFT           = Alignment(horizontal="left",   vertical="center")

LOGO_PATH = os.path.join(os.path.dirname(__file__), "LP Logo.png")

# Col A is a blank spacer — data occupies cols B (2) through J (10)
COL_A_WIDTH = 4

COLS = {
    "campaign":    2,   # B
    "impressions": 3,   # C
    "clicks":      4,   # D
    "cost":        5,   # E
    "ctr":         6,   # F
    "conversions": 7,   # G
    "conv_value":  8,   # H
    "cpa":         9,   # I
    "roas":        10,  # J
}
COL_HEADERS = ["Campaign", "Impressions", "Clicks", "Cost", "CTR %",
               "Conversions", "Conv Value", "CPA", "ROAS"]
COL_WIDTHS  = [40, 14, 10, 12, 10, 14, 14, 12, 10]   # widths for cols B–J
COL_FMTS    = [None, "#,##0", "#,##0", "[$\xa3-809]#,##0.00", "0.00%",
               "#,##0.00", "[$\xa3-809]#,##0.00",
               "[$\xa3-809]#,##0.00", '0.00"x"']

# Data tab column positions (1-based)
D_DATE  = 1
D_CAMP  = 2
D_IMPR  = 3
D_CLICK = 4
D_COST  = 5
D_CONV  = 6
D_CVAL  = 7

# Search term report styles (reuse body/header styles)
_BORDER_THIN = Border(bottom=Side(style="thin", color=BORDER_CLR))


def _client_path(base: str, account_name: str, folder: str, folder_name: str = None) -> str:
    path = os.path.join(base, folder_name or account_name, folder)
    os.makedirs(path, exist_ok=True)
    return path


def _style(cell, fill, font, alignment, border=_BORDER_THIN, fmt=None):
    cell.fill      = fill
    cell.font      = font
    cell.alignment = alignment
    cell.border    = border
    if fmt:
        cell.number_format = fmt


def _box_border(ws, min_row, max_row, min_col, max_col):
    """Apply a black medium box border around the outer edges of a cell range."""
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            cell = ws.cell(r, c)
            top    = _BLACK_MEDIUM if r == min_row else cell.border.top
            bottom = _BLACK_MEDIUM if r == max_row else cell.border.bottom
            left   = _BLACK_MEDIUM if c == min_col else cell.border.left
            right  = _BLACK_MEDIUM if c == max_col else cell.border.right
            cell.border = Border(top=top, bottom=bottom, left=left, right=right)


def _header_row(ws, row_num, heights=22):
    """Write amber header row; col A is white spacer, cols B-J get column labels."""
    ws.row_dimensions[row_num].height = heights
    ws.cell(row_num, 1, "")
    _style(ws.cell(row_num, 1), _WHITE_FILL, _BODY_FONT, _CENTER, border=_NO_BORDER)
    for col, (label, fmt) in enumerate(zip(COL_HEADERS, COL_FMTS), 2):
        c = ws.cell(row_num, col, label)
        _style(c, _HEADER_FILL, _HEADER_FONT, _CENTER, border=_NO_BORDER, fmt=fmt)


def _sumifs(data_col, p_start_row, p_end_row, campaign):
    """SUMIFS formula: sum data_col where date in period and campaign matches."""
    dc = get_column_letter(data_col)
    return (
        f"=SUMIFS(Data!${dc}:${dc},"
        f"Data!${get_column_letter(D_DATE)}:${get_column_letter(D_DATE)},"
        f'">="&$P${p_start_row},'
        f"Data!${get_column_letter(D_DATE)}:${get_column_letter(D_DATE)},"
        f'"<="&$P${p_end_row},'
        f"Data!${get_column_letter(D_CAMP)}:${get_column_letter(D_CAMP)},"
        f'"{campaign}")'
    )


def _write_campaign_rows(ws, campaigns, start_row, p_start_row, p_end_row):
    """Write per-campaign rows with SUMIFS formulas. Returns last row written."""
    row = start_row
    for camp in campaigns:
        fill = _ALT_FILL if (row - start_row) % 2 == 0 else _WHITE_FILL
        ws.row_dimensions[row].height = 18

        # A: blank spacer
        _style(ws.cell(row, 1, ""), fill, _BODY_FONT, _CENTER, border=_NO_BORDER)

        # B: Campaign name
        _style(ws.cell(row, COLS["campaign"], camp), fill, _BODY_FONT, _CENTER, border=_NO_BORDER)

        # C: Impressions
        c = ws.cell(row, COLS["impressions"])
        c.value = _sumifs(D_IMPR, p_start_row, p_end_row, camp)
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt="#,##0")

        # D: Clicks
        c = ws.cell(row, COLS["clicks"])
        c.value = _sumifs(D_CLICK, p_start_row, p_end_row, camp)
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt="#,##0")

        # E: Cost
        c = ws.cell(row, COLS["cost"])
        c.value = _sumifs(D_COST, p_start_row, p_end_row, camp)
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt='[$\xa3-809]#,##0.00')

        # F: CTR % = Clicks / Impressions  (D/C)
        c = ws.cell(row, COLS["ctr"])
        c.value = f"=IF(C{row}>0,D{row}/C{row},0)"
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt="0.00%")

        # G: Conversions
        c = ws.cell(row, COLS["conversions"])
        c.value = _sumifs(D_CONV, p_start_row, p_end_row, camp)
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt="#,##0.00")

        # H: Conv Value
        c = ws.cell(row, COLS["conv_value"])
        c.value = _sumifs(D_CVAL, p_start_row, p_end_row, camp)
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt='[$\xa3-809]#,##0.00')

        # I: CPA = Cost / Conversions  (E/G)
        c = ws.cell(row, COLS["cpa"])
        c.value = f'=IF(G{row}>0,E{row}/G{row},"")'
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt='[$\xa3-809]#,##0.00')

        # J: ROAS = Conv Value / Cost  (H/E)
        c = ws.cell(row, COLS["roas"])
        c.value = f'=IF(E{row}>0,H{row}/E{row},"")'
        _style(c, fill, _BODY_FONT, _CENTER, border=_NO_BORDER, fmt='0.00"x"')

        row += 1
    return row - 1


def _write_total_row(ws, row, label, first_data, last_data, fill=None, font=None):
    fill = fill or _HEADER_FILL
    font = font or _HEADER_FONT
    ws.row_dimensions[row].height = 22

    # A: blank spacer (always white)
    _style(ws.cell(row, 1, ""), _WHITE_FILL, _BODY_FONT, _CENTER, border=_NO_BORDER)

    # B: label
    _style(ws.cell(row, COLS["campaign"], label), fill, font, _CENTER, border=_NO_BORDER)

    sums = [
        (COLS["impressions"], f"=SUM(C{first_data}:C{last_data})", "#,##0"),
        (COLS["clicks"],      f"=SUM(D{first_data}:D{last_data})", "#,##0"),
        (COLS["cost"],        f"=SUM(E{first_data}:E{last_data})", '[$\xa3-809]#,##0.00'),
        (COLS["ctr"],         f"=IF(C{row}>0,D{row}/C{row},0)",    "0.00%"),
        (COLS["conversions"], f"=SUM(G{first_data}:G{last_data})", "#,##0.00"),
        (COLS["conv_value"],  f"=SUM(H{first_data}:H{last_data})", '[$\xa3-809]#,##0.00'),
        (COLS["cpa"],         f'=IF(G{row}>0,E{row}/G{row},"")',   '[$\xa3-809]#,##0.00'),
        (COLS["roas"],        f'=IF(E{row}>0,H{row}/E{row},"")',   '0.00"x"'),
    ]
    for col, formula, fmt in sums:
        c = ws.cell(row, col)
        c.value = formula
        _style(c, fill, font, _CENTER, border=_NO_BORDER, fmt=fmt)


def _write_pct_change_row(ws, row, total_row, pct_values=None):
    """
    Write a % change row (current vs previous 7 days).
    Formulas pull from SUMIFS on the Data tab (P3/P4 anchors).
    pct_values dict (pre-computed from raw data) drives static green/red font colours,
    which works in Numbers and Excel alike (no conditional formatting rules needed).
    Format: +15.2% / -8.3% / – (zero)
    """
    ws.row_dimensions[row].height = 20

    dd = get_column_letter(D_DATE)

    def _prev(data_col):
        dc = get_column_letter(data_col)
        return (
            f"SUMIFS(Data!${dc}:${dc},"
            f"Data!${dd}:${dd},"
            f'">="&$P$3,'
            f"Data!${dd}:${dd},"
            f'"<="&$P$4)'
        )

    pi = _prev(D_IMPR)
    pk = _prev(D_CLICK)
    pe = _prev(D_COST)
    pg = _prev(D_CONV)
    ph = _prev(D_CVAL)
    T  = total_row
    pct_fmt = '+0.0%;-0.0%;"–"'

    _style(ws.cell(row, 1, ""), _WHITE_FILL, _BODY_FONT, _CENTER, border=_NO_BORDER)
    _style(ws.cell(row, COLS["campaign"], "vs prev 7 days"),
           _WHITE_FILL, _BOLD_DARK_FONT, _CENTER, border=_NO_BORDER)

    # positive = good for these; negative = good for cost & cpa
    colour_map = {
        COLS["impressions"]: ("impressions", True),
        COLS["clicks"]:      ("clicks",      True),
        COLS["cost"]:        ("cost",        False),
        COLS["ctr"]:         ("ctr",         True),
        COLS["conversions"]: ("conversions", True),
        COLS["conv_value"]:  ("conv_value",  True),
        COLS["cpa"]:         ("cpa",         False),
        COLS["roas"]:        ("roas",        True),
    }

    formulas = [
        (COLS["impressions"], f'=IF({pi}>0,(C{T}/{pi})-1,"")'),
        (COLS["clicks"],      f'=IF({pk}>0,(D{T}/{pk})-1,"")'),
        (COLS["cost"],        f'=IF({pe}>0,(E{T}/{pe})-1,"")'),
        (COLS["ctr"],         f'=IF({pi}>0,IF({pk}/{pi}>0,(F{T}/({pk}/{pi}))-1,""),"")'),
        (COLS["conversions"], f'=IF({pg}>0,(G{T}/{pg})-1,"")'),
        (COLS["conv_value"],  f'=IF({ph}>0,(H{T}/{ph})-1,"")'),
        (COLS["cpa"],         f'=IF({pg}>0,IF({pe}/{pg}>0,(I{T}/({pe}/{pg}))-1,""),"")'),
        (COLS["roas"],        f'=IF({pe}>0,IF({ph}/{pe}>0,(J{T}/({ph}/{pe}))-1,""),"")'),
    ]
    for col, formula in formulas:
        c = ws.cell(row, col)
        c.value = formula
        font = _BODY_FONT
        if pct_values:
            key, pos_good = colour_map[col]
            val = pct_values.get(key)
            if val is not None and abs(val) > 0.001:
                font = _CF_GREEN if ((val > 0) == pos_good) else _CF_RED
        _style(c, _WHITE_FILL, font, _CENTER, border=_NO_BORDER, fmt=pct_fmt)


def _generate_analysis(account_name, curr, pct):
    """Call Claude Haiku for a 2-3 sentence client-facing performance summary.
    Falls back to a formulaic sentence if the API is unavailable."""
    try:
        import anthropic
        api_key = os.environ.get("ANTHROPIC_API_KEY", "")
        if not api_key:
            raise ValueError("no key")

        def _fmt(val, suffix=""):
            return f"{val:+.1f}{suffix}" if val is not None else "N/A"

        prompt = (
            f"You are a PPC account manager writing a brief weekly update for {account_name}.\n\n"
            f"This week's results:\n"
            f"- Spend: £{curr['cost']:,.2f} ({_fmt(pct.get('cost'), '%')} vs last week)\n"
            f"- Clicks: {curr['clicks']:,} ({_fmt(pct.get('clicks'), '%')} vs last week)\n"
            f"- Impressions: {curr['impressions']:,} ({_fmt(pct.get('impressions'), '%')} vs last week)\n"
            f"- Conversions: {curr['conversions']:,.1f} ({_fmt(pct.get('conversions'), '%')} vs last week)\n"
        )
        if curr["cpa"] is not None:
            prompt += f"- CPA: £{curr['cpa']:,.2f} ({_fmt(pct.get('cpa'), '%')} vs last week)\n"
        if curr["roas"] is not None:
            prompt += f"- ROAS: {curr['roas']:.2f}x ({_fmt(pct.get('roas'), '%')} vs last week)\n"
        prompt += (
            "\nWrite 2-3 sentences for the client. Be direct. Lead with the headline result. "
            "Note whether efficiency (CPA/ROAS) improved or worsened. No jargon, no bullet points."
        )

        ai_client = anthropic.Anthropic(api_key=api_key)
        msg = ai_client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=200,
            messages=[{"role": "user", "content": prompt}],
        )
        return msg.content[0].text.strip()
    except Exception:
        # Formulaic fallback
        parts = []
        if pct.get("conversions") is not None:
            direction = "up" if pct["conversions"] > 0 else "down"
            parts.append(f"Conversions were {direction} {abs(pct['conversions']):.1f}% week-on-week.")
        if curr.get("cpa") is not None and pct.get("cpa") is not None:
            direction = "improved" if pct["cpa"] < 0 else "increased"
            parts.append(f"CPA {direction} to £{curr['cpa']:,.2f} ({pct['cpa']:+.1f}%).")
        elif curr.get("roas") is not None and pct.get("roas") is not None:
            direction = "improved" if pct["roas"] > 0 else "declined"
            parts.append(f"ROAS {direction} to {curr['roas']:.2f}x ({pct['roas']:+.1f}%).")
        return " ".join(parts) if parts else f"Weekly report for {account_name}."


# ── Shared xlsx builder ───────────────────────────────────────────────────────

def _build_perf_workbook(daily_rows, current_start, current_end, prev_start, prev_end):
    """Build the branded performance workbook. Returns (wb, curr, pct, camp_agg, campaigns)."""
    curr_rows = [r for r in daily_rows if current_start.isoformat() <= r["date"] <= current_end.isoformat()]
    prev_rows = [r for r in daily_rows if prev_start.isoformat()    <= r["date"] <= prev_end.isoformat()]

    def _agg(rows):
        cost   = sum(r["cost"]        for r in rows)
        clicks = sum(r["clicks"]      for r in rows)
        impr   = sum(r["impressions"] for r in rows)
        conv   = sum(r["conversions"] for r in rows)
        cval   = sum(r["conv_value"]  for r in rows)
        return {
            "cost": cost, "clicks": clicks, "impressions": impr,
            "conversions": conv, "conv_value": cval,
            "ctr":  clicks / impr if impr > 0 else None,
            "cpa":  cost   / conv if conv > 0 else None,
            "roas": cval   / cost if cost > 0 else None,
        }

    curr = _agg(curr_rows)
    prev = _agg(prev_rows)

    def _pct(c, p):
        return ((c - p) / p * 100) if p and p > 0 else None

    pct = {
        "cost":        _pct(curr["cost"],        prev["cost"]),
        "clicks":      _pct(curr["clicks"],      prev["clicks"]),
        "impressions": _pct(curr["impressions"], prev["impressions"]),
        "conversions": _pct(curr["conversions"], prev["conversions"]),
        "conv_value":  _pct(curr["conv_value"],  prev["conv_value"]),
        "ctr":         _pct(curr["ctr"]  or 0,   prev["ctr"])  if prev["ctr"]  else None,
        "cpa":         _pct(curr["cpa"]  or 0,   prev["cpa"])  if prev["cpa"]  else None,
        "roas":        _pct(curr["roas"] or 0,   prev["roas"]) if prev["roas"] else None,
    }

    camp_agg = defaultdict(lambda: {"cost": 0, "clicks": 0, "impressions": 0,
                                     "conversions": 0, "conv_value": 0})
    for r in curr_rows:
        d = camp_agg[r["campaign_name"]]
        for k in ("cost", "clicks", "impressions", "conversions", "conv_value"):
            d[k] += r[k]
    campaigns = sorted(camp_agg, key=lambda c: -camp_agg[c]["cost"])

    wb   = Workbook()
    ws_r = wb.active
    ws_r.title = "Results"
    ws_r.sheet_view.showGridLines = False

    for _r in range(1, 55):
        for _c in range(1, 18):
            cell = ws_r.cell(_r, _c)
            cell.border = _NO_BORDER
            cell.fill   = _WHITE_FILL

    for p_row, val in [(1, current_start.isoformat()),
                       (2, current_end.isoformat()),
                       (3, prev_start.isoformat()),
                       (4, prev_end.isoformat())]:
        ws_r.cell(p_row, 16, val)
    ws_r.column_dimensions["P"].hidden = True

    ws_r.row_dimensions[1].height = 15
    ws_r.row_dimensions[2].height = 45
    ws_r.row_dimensions[3].height = 45

    c = ws_r.cell(2, 2, f"Weekly Report: {current_start.strftime('%d %b')} – {current_end.strftime('%d %b %Y')}")
    _style(c, _WHITE_FILL, _TITLE_FONT, _CENTER, border=_NO_BORDER)

    if os.path.exists(LOGO_PATH):
        logo = XLImage(LOGO_PATH)
        logo.width  = 80
        logo.height = 80
        ws_r.add_image(logo, "F2")

    ws_r.column_dimensions["A"].width = COL_A_WIDTH
    for col_idx, width in enumerate(COL_WIDTHS, 2):
        ws_r.column_dimensions[get_column_letter(col_idx)].width = width

    HEADER_ROW     = 4
    FIRST_DATA_ROW = HEADER_ROW + 1

    _header_row(ws_r, HEADER_ROW)
    last_camp_row = _write_campaign_rows(ws_r, campaigns, FIRST_DATA_ROW, 1, 2)

    TOTAL_ROW = last_camp_row + 2
    _write_total_row(ws_r, TOTAL_ROW, "Total Performance", FIRST_DATA_ROW, last_camp_row)

    _box_border(ws_r, HEADER_ROW, last_camp_row, 2, 10)
    _box_border(ws_r, TOTAL_ROW, TOTAL_ROW, 2, 10)

    PCT_ROW = TOTAL_ROW + 1
    _write_pct_change_row(ws_r, PCT_ROW, TOTAL_ROW, pct_values=pct)
    _box_border(ws_r, PCT_ROW, PCT_ROW, 2, 10)

    ws_d = wb.create_sheet("Data")
    ws_d.freeze_panes = "A2"

    data_headers = ["Date", "Campaign", "Impressions", "Clicks", "Cost", "Conversions", "Conv Value"]
    data_widths  = [14, 40, 14, 10, 14, 14, 14]
    data_fmts    = [None, None, "#,##0", "#,##0", '[$\xa3-809]#,##0.00', "#,##0.00", '[$\xa3-809]#,##0.00']

    ws_d.row_dimensions[1].height = 22
    for col_idx, (header, width) in enumerate(zip(data_headers, data_widths), 1):
        c = ws_d.cell(1, col_idx, header)
        _style(c, _HEADER_FILL, _HEADER_FONT, _CENTER)
        ws_d.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, row in enumerate(daily_rows, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else _WHITE_FILL
        ws_d.row_dimensions[row_idx].height = 18
        for col_idx, (key, fmt) in enumerate(zip(
            ["date", "campaign_name", "impressions", "clicks", "cost", "conversions", "conv_value"],
            data_fmts
        ), 1):
            c = ws_d.cell(row_idx, col_idx, row[key])
            _style(c, fill, _BODY_FONT, _LEFT, fmt=fmt)

    return wb, curr, pct, camp_agg, campaigns


# ── Public writers ────────────────────────────────────────────────────────────

def write_performance_report(
    base: str,
    account_name: str,
    daily_rows: list,
    current_start: date,
    current_end: date,
    prev_start: date,
    prev_end: date,
    folder_name: str = None,
) -> str:
    """Write a two-tab branded xlsx (Results + Data) plus Obsidian markdown. Returns the markdown path."""
    folder        = _client_path(base, account_name, "Performance Reports", folder_name)
    date_slug     = current_end.strftime("%Y-%m-%d")
    md_filename   = f"performance_{date_slug}.md"
    xlsx_filename = f"performance_{date_slug}.xlsx"
    md_path       = os.path.join(folder, md_filename)
    xlsx_path     = os.path.join(folder, xlsx_filename)

    wb, curr, pct, camp_agg, campaigns = _build_perf_workbook(
        daily_rows, current_start, current_end, prev_start, prev_end
    )
    wb.save(xlsx_path)

    def _arrow(val, good_when_positive=True):
        if val is None:
            return "—"
        symbol = "↑" if val > 0 else ("↓" if val < 0 else "→")
        good   = (val > 0) == good_when_positive
        tag    = " ✅" if good else " ⚠️"
        return f"{val:+.1f}% {symbol}{tag}"

    analysis = _generate_analysis(account_name, curr, pct)

    lines = [
        "---",
        f"tags: [google-ads, performance-report, {account_name.lower().replace(' ', '-')}]",
        f"date: {date.today().isoformat()}",
        f"account: {account_name}",
        f"period: \"{current_start} → {current_end}\"",
        "---",
        "",
        f"# {account_name} — Weekly Performance Report",
        f"**Period:** {current_start.strftime('%d %b %Y')} → {current_end.strftime('%d %b %Y')}",
        "",
        "> [!note] Spreadsheet",
        f"> [[{xlsx_filename}|⬇ Open in Numbers]]",
        "",
        "## Performance Summary",
        "",
        "| Metric | This Week | vs Last Week |",
        "| --- | ---: | --- |",
        f"| Spend | £{curr['cost']:,.2f} | {_arrow(pct.get('cost'), good_when_positive=False)} |",
        f"| Clicks | {curr['clicks']:,} | {_arrow(pct.get('clicks'))} |",
        f"| Impressions | {curr['impressions']:,} | {_arrow(pct.get('impressions'))} |",
        f"| Conversions | {curr['conversions']:,.1f} | {_arrow(pct.get('conversions'))} |",
        f"| Conv. Value | £{curr['conv_value']:,.2f} | {_arrow(pct.get('conv_value'))} |",
    ]
    if curr["cpa"] is not None:
        lines.append(f"| CPA | £{curr['cpa']:,.2f} | {_arrow(pct.get('cpa'), good_when_positive=False)} |")
    if curr["roas"] is not None:
        lines.append(f"| ROAS | {curr['roas']:.2f}x | {_arrow(pct.get('roas'))} |")

    lines += ["", "## Campaign Breakdown", "",
              "| Campaign | Spend | Clicks | Conv | CPA | ROAS |",
              "| --- | ---: | ---: | ---: | ---: | ---: |"]
    for camp in campaigns:
        d    = camp_agg[camp]
        cpa  = f"£{d['cost']/d['conversions']:,.2f}" if d["conversions"] > 0 else "—"
        roas = f"{d['conv_value']/d['cost']:.2f}x"   if d["cost"]        > 0 else "—"
        lines.append(
            f"| {camp} | £{d['cost']:,.2f} | {d['clicks']:,} "
            f"| {d['conversions']:,.1f} | {cpa} | {roas} |"
        )

    lines += ["", "## Analysis", "", analysis, ""]

    with open(md_path, "w") as f:
        f.write("\n".join(lines) + "\n")

    return md_path


def write_weekly_report(
    base: str,
    account_name: str,
    daily_rows: list,
    current_start: date,
    current_end: date,
    prev_start: date,
    prev_end: date,
    irrelevant_terms: list,
    date_from_str: str,
    date_to_str: str,
    deep_dive_md: str,
    folder_name: str = None,
) -> str:
    """
    Write a combined weekly report: performance xlsx + optional search terms xlsx +
    single Obsidian note containing all three sections. Returns the markdown path.
    """
    folder    = _client_path(base, account_name, "Performance Reports", folder_name)
    date_slug = current_end.strftime("%Y-%m-%d")

    perf_xlsx  = f"performance_{date_slug}.xlsx"
    terms_xlsx = f"irrelevant_terms_{date_slug}.xlsx"
    md_path    = os.path.join(folder, f"weekly_report_{date_slug}.md")

    # ── Performance xlsx ──────────────────────────────────────────────────────
    wb, curr, pct, camp_agg, campaigns = _build_perf_workbook(
        daily_rows, current_start, current_end, prev_start, prev_end
    )
    wb.save(os.path.join(folder, perf_xlsx))

    # ── Search terms xlsx (only if there are irrelevant terms) ─────────────
    has_terms = bool(irrelevant_terms)
    if has_terms:
        st_headers = ["Negative Keyword", "Match Type", "Campaign Name", "Ad Group Name",
                      "Impressions", "Clicks", "Cost", "CTR %",
                      "Conversions", "Conv Value", "Reason"]
        st_widths  = [35, 14, 38, 24, 14, 10, 12, 10, 14, 14, 40]

        wb_st = Workbook()
        ws_st = wb_st.active
        ws_st.title = "Irrelevant Terms"
        ws_st.freeze_panes = "A2"
        ws_st.row_dimensions[1].height = 22

        for col_idx, (header, width) in enumerate(zip(st_headers, st_widths), 1):
            c = ws_st.cell(1, col_idx, header)
            _style(c, _HEADER_FILL, _HEADER_FONT, _CENTER)
            ws_st.column_dimensions[get_column_letter(col_idx)].width = width

        for row_idx, r in enumerate(irrelevant_terms, 2):
            fill = _ALT_FILL if row_idx % 2 == 0 else _WHITE_FILL
            ws_st.row_dimensions[row_idx].height = 18
            values = [
                r["search_term"], "Exact Match", r["campaign_name"], r["ad_group_name"],
                r["impressions"], r["clicks"], f"£{r['cost']:,.2f}", f"{r['ctr_pct']}%",
                r["conversions"], f"£{r['conv_value']:,.2f}", r.get("reason", ""),
            ]
            for col_idx, val in enumerate(values, 1):
                c = ws_st.cell(row_idx, col_idx, val)
                _style(c, fill, _BODY_FONT, _CENTER)

        wb_st.save(os.path.join(folder, terms_xlsx))

    # ── Shared markdown helpers ───────────────────────────────────────────────
    def _arrow(val, good_when_positive=True):
        if val is None:
            return "—"
        symbol = "↑" if val > 0 else ("↓" if val < 0 else "→")
        good   = (val > 0) == good_when_positive
        tag    = " ✅" if good else " ⚠️"
        return f"{val:+.1f}% {symbol}{tag}"

    analysis = _generate_analysis(account_name, curr, pct)

    # ── Build combined markdown ───────────────────────────────────────────────
    tag_slug = account_name.lower().replace(" ", "-")
    links    = f"[[{perf_xlsx}|📊 Performance Data]]"
    if has_terms:
        links += f"  ·  [[{terms_xlsx}|🔍 Search Terms]]"

    lines = [
        "---",
        f"tags: [google-ads, weekly-report, {tag_slug}]",
        f"date: {date.today().isoformat()}",
        f"account: {account_name}",
        f"period: \"{current_start} → {current_end}\"",
        "---",
        "",
        f"# {account_name} — Weekly Report",
        f"**Period:** {current_start.strftime('%d %b %Y')} → {current_end.strftime('%d %b %Y')}",
        "",
        f"> {links}",
        "",
        "---",
        "",
        "## Performance Summary",
        "",
        "| Metric | This Week | vs Last Week |",
        "| --- | ---: | --- |",
        f"| Spend | £{curr['cost']:,.2f} | {_arrow(pct.get('cost'), good_when_positive=False)} |",
        f"| Clicks | {curr['clicks']:,} | {_arrow(pct.get('clicks'))} |",
        f"| Impressions | {curr['impressions']:,} | {_arrow(pct.get('impressions'))} |",
        f"| Conversions | {curr['conversions']:,.1f} | {_arrow(pct.get('conversions'))} |",
        f"| Conv. Value | £{curr['conv_value']:,.2f} | {_arrow(pct.get('conv_value'))} |",
    ]
    if curr["cpa"] is not None:
        lines.append(f"| CPA | £{curr['cpa']:,.2f} | {_arrow(pct.get('cpa'), good_when_positive=False)} |")
    if curr["roas"] is not None:
        lines.append(f"| ROAS | {curr['roas']:.2f}x | {_arrow(pct.get('roas'))} |")

    lines += ["", "## Campaign Breakdown", "",
              "| Campaign | Spend | Clicks | Conv | CPA | ROAS |",
              "| --- | ---: | ---: | ---: | ---: | ---: |"]
    for camp in campaigns:
        d    = camp_agg[camp]
        cpa  = f"£{d['cost']/d['conversions']:,.2f}" if d["conversions"] > 0 else "—"
        roas = f"{d['conv_value']/d['cost']:.2f}x"   if d["cost"]        > 0 else "—"
        lines.append(
            f"| {camp} | £{d['cost']:,.2f} | {d['clicks']:,} "
            f"| {d['conversions']:,.1f} | {cpa} | {roas} |"
        )

    lines += ["", "## Client Summary", "", analysis, "", "---", ""]

    # ── Search Term Audit section ─────────────────────────────────────────────
    if has_terms:
        total_spend  = sum(r["cost"]   for r in irrelevant_terms)
        total_clicks = sum(r["clicks"] for r in irrelevant_terms)
        lines += [
            "## Search Term Audit",
            "",
            f"**Irrelevant terms found:** {len(irrelevant_terms)}  ",
            f"**Estimated wasted spend:** £{total_spend:,.2f}  ",
            f"**Wasted clicks:** {total_clicks:,}",
            "",
            "| Search Term | Ad Group | Clicks | Cost | Why Irrelevant |",
            "| --- | --- | ---: | ---: | --- |",
        ]
        for r in irrelevant_terms:
            reason = r.get("reason", "Off-intent").replace("|", "/")
            lines.append(
                f"| {r['search_term'].replace('|','/')} "
                f"| {r['ad_group_name'].replace('|','/')} "
                f"| {r['clicks']:,} | £{r['cost']:,.2f} | {reason} |"
            )
        unique_terms = sorted({r["search_term"] for r in irrelevant_terms})
        lines += ["", "### Recommended Negative Keywords", "",
                  "> Add as **exact match** negatives at campaign or account level.", ""]
        for t in unique_terms:
            lines.append(f"- [ ] `[{t}]`")
        lines += ["", "---", ""]
    else:
        lines += [
            "## Search Term Audit",
            "",
            "*No irrelevant terms found this week.*",
            "",
            "---",
            "",
        ]

    # ── Account Manager Deep Dive section ─────────────────────────────────────
    lines += [
        "## Account Manager Deep Dive",
        "",
        f"> [!tip] Internal use only",
        "> Work through the Priority Actions before preparing the client update.",
        "",
        deep_dive_md,
        "",
    ]

    with open(md_path, "w") as f:
        f.write("\n".join(lines) + "\n")

    return md_path


def write_search_term_report(
    base: str,
    account_name: str,
    date_from: str,
    date_to: str,
    irrelevant_terms: list,
    folder_name: str = None,
) -> str:
    """Write a search term irrelevance report markdown + xlsx into the Obsidian vault."""
    folder        = _client_path(base, account_name, "Search Term Reports", folder_name)
    date_slug     = date.today().strftime("%Y-%m-%d")
    md_filename   = f"irrelevant_terms_{date_slug}.md"
    xlsx_filename = f"irrelevant_terms_{date_slug}.xlsx"
    filepath      = os.path.join(folder, md_filename)
    xlsx_path     = os.path.join(folder, xlsx_filename)

    total_spend  = sum(r["cost"] for r in irrelevant_terms)
    total_clicks = sum(r["clicks"] for r in irrelevant_terms)

    st_headers = ["Negative Keyword", "Match Type", "Campaign Name", "Ad Group Name",
                  "Impressions", "Clicks", "Cost", "CTR %",
                  "Conversions", "Conv Value", "Reason"]
    st_widths  = [35, 14, 38, 24, 14, 10, 12, 10, 14, 14, 40]

    wb = Workbook()
    ws = wb.active
    ws.title = "Irrelevant Terms"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for col_idx, (header, width) in enumerate(zip(st_headers, st_widths), 1):
        c = ws.cell(1, col_idx, header)
        _style(c, _HEADER_FILL, _HEADER_FONT, _CENTER)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, r in enumerate(irrelevant_terms, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else _WHITE_FILL
        ws.row_dimensions[row_idx].height = 18
        values = [
            r["search_term"], "Exact Match", r["campaign_name"], r["ad_group_name"],
            r["impressions"], r["clicks"], f"£{r['cost']:,.2f}", f"{r['ctr_pct']}%",
            r["conversions"], f"£{r['conv_value']:,.2f}", r.get("reason", ""),
        ]
        for col_idx, val in enumerate(values, 1):
            c = ws.cell(row_idx, col_idx, val)
            _style(c, fill, _BODY_FONT, _CENTER)

    wb.save(xlsx_path)

    lines = [
        "---",
        f"tags: [google-ads, search-terms, negatives, {account_name.lower().replace(' ', '-')}]",
        f"date: {date.today().isoformat()}",
        f"account: {account_name}",
        f"period: \"{date_from} → {date_to}\"",
        "---",
        "",
        f"# {account_name} — Irrelevant Search Terms",
        f"**Period:** {date_from} → {date_to}  ",
        f"**Terms found:** {len(irrelevant_terms)}  ",
        f"**Estimated wasted spend:** £{total_spend:,.2f}  ",
        f"**Wasted clicks:** {total_clicks:,}",
        "",
        "> [!note] Raw data",
        f"> [[{xlsx_filename}|⬇ Download spreadsheet]]",
        "",
        "## Irrelevant Terms",
        "",
        "| Negative Keyword | Match Type | Ad Group | Clicks | Cost | Conv Value | Why Irrelevant |",
        "| --- | --- | --- | ---: | ---: | ---: | --- |",
    ]
    for r in irrelevant_terms:
        reason = r.get("reason", "Off-intent for account theme").replace("|", "/")
        lines.append(
            f"| {r['search_term'].replace('|','/')} | Exact Match "
            f"| {r['ad_group_name'].replace('|','/')} | {r['clicks']:,} "
            f"| £{r['cost']:,.2f} | £{r['conv_value']:,.2f} | {reason} |"
        )

    unique_terms = sorted({r["search_term"] for r in irrelevant_terms})
    lines += ["", "## Recommended Negative Keywords", "",
              "> Add these as **exact match** negatives at campaign or account level.", ""]
    for t in unique_terms:
        lines.append(f"- [ ] `[{t}]`")

    with open(filepath, "w") as f:
        f.write("\n".join(lines) + "\n")

    return filepath
